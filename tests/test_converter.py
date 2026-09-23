# -*- coding: utf-8 -*-
"""PDFToWordConverter 单元测试 + 端到端测试

运行: python -X utf8 -m unittest discover -s tests -v
"""

import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pdf_to_word_converter import PDFToWordConverter

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(PROJECT_DIR, "data")


def _make_sample_pdf(path, text="Hello PDF to Word 转换测试\n第二行内容"):
    """用 PyMuPDF 生成一个含文本层的简单 PDF"""
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), text, fontsize=14)
    page.insert_text((72, 140), "Table-like line: col1  col2  col3", fontsize=12)
    page2 = doc.new_page()
    page2.insert_text((72, 100), "Page two content", fontsize=14)
    doc.save(path)
    doc.close()


def _make_mixed_pdf(path):
    """生成混合 PDF: 第1页文本层, 第2~3页纯图片(含文字), 第4页纯色照片页"""
    import fitz
    doc = fitz.open()

    # 第1页: 真实文本层(文字页)
    page = doc.new_page(width=595, height=842)
    page.insert_text((72, 100), "QUANTUM COVER PAGE WITH REAL TEXT", fontsize=14)

    # 第2~3页: 文本渲染成位图后整页插入 => 扫描图片页
    for i, label in enumerate(("ALPHA", "BETA"), start=2):
        tmp = fitz.open()
        tp = tmp.new_page(width=595, height=842)
        tp.insert_text((72, 120), f"HYBRID OCR PAGE {label}", fontsize=18)
        pix = tp.get_pixmap(dpi=100)
        tmp.close()
        page = doc.new_page(width=595, height=842)
        page.insert_image(page.rect, pixmap=pix)

    # 第4页: 纯色图(无文字) => 应触发"嵌入原图"兜底
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 595, 842))
    pix.set_pixel(0, 0, (200, 200, 200))
    pix.clear_with(180)
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, pixmap=pix)

    doc.save(path)
    doc.close()


class TestOutputFilename(unittest.TestCase):
    def test_docx_suffix(self):
        c = PDFToWordConverter('in', 'out')
        self.assertEqual(c._generate_output_filename('报告.pdf'), '报告.docx')
        self.assertEqual(c._generate_output_filename('a.b.c.pdf'), 'a.b.c.docx')


class TestResultShape(unittest.TestCase):
    def test_result_keys(self):
        r = PDFToWordConverter._result('t.pdf', 'error', error='x')
        for key in ('file', 'status', 'pages', 'output', 'warning', 'error', 'elapsed'):
            self.assertIn(key, r)
        self.assertEqual(r['status'], 'error')


class TestConvertSingle(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="p2w_test_")
        self.in_dir = os.path.join(self.tmp, "in")
        self.out_dir = os.path.join(self.tmp, "out")
        os.makedirs(self.in_dir)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_missing_file_error(self):
        c = PDFToWordConverter(self.in_dir, self.out_dir)
        r = c._convert_single_pdf(os.path.join(self.in_dir, "nope.pdf"))
        self.assertEqual(r['status'], 'error')

    def test_corrupt_file_error(self):
        bad = os.path.join(self.in_dir, "bad.pdf")
        with open(bad, "wb") as f:
            f.write(b"not a pdf at all")
        c = PDFToWordConverter(self.in_dir, self.out_dir)
        r = c._convert_single_pdf(bad)
        self.assertEqual(r['status'], 'error')
        self.assertTrue(r['error'])

    def test_end_to_end_generated_pdf(self):
        pdf = os.path.join(self.in_dir, "sample.pdf")
        _make_sample_pdf(pdf)
        c = PDFToWordConverter(self.in_dir, self.out_dir)
        results = c.convert_files([pdf])
        self.assertEqual(len(results), 1)
        r = results[0]
        self.assertEqual(r['status'], 'ok', r.get('error'))
        self.assertEqual(r['pages'], 2)
        self.assertEqual(r['output'], 'sample.docx')
        out_path = os.path.join(self.out_dir, r['output'])
        self.assertTrue(os.path.isfile(out_path))
        self.assertGreater(os.path.getsize(out_path), 0)

        # docx 可读且含正文
        from docx import Document
        doc = Document(out_path)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Hello PDF to Word", text)
        self.assertIn("Page two", text)

    def test_batch_name_collision(self):
        pdf1 = os.path.join(self.in_dir, "a.pdf")
        _make_sample_pdf(pdf1, "file one")
        c = PDFToWordConverter(self.in_dir, self.out_dir)
        # 模拟批内同名: 手动构造 used_names 已占用 a.docx
        used = {"a.docx"}
        r = c._convert_single_pdf(pdf1, used)
        self.assertEqual(r['status'], 'ok', r.get('error'))
        self.assertEqual(r['output'], 'a(2).docx')
        self.assertIn('a(2).docx', used)

    def test_on_progress_callback(self):
        pdf = os.path.join(self.in_dir, "sample.pdf")
        _make_sample_pdf(pdf)
        c = PDFToWordConverter(self.in_dir, self.out_dir)
        calls = []
        c.convert_files([pdf], on_progress=lambda res, i, n: calls.append((i, n)))
        self.assertEqual(calls, [(1, 1)])

    def test_ocr_page_progress_callback(self):
        trump = os.path.join(PROJECT_DIR, "特朗普.pdf")
        if not os.path.isfile(trump):
            self.skipTest("特朗普.pdf 不存在")
        c = PDFToWordConverter(self.in_dir, self.out_dir)
        pages_seen = []
        rs = c.convert_files(
            [trump],
            on_page_progress=lambda cur, tot: pages_seen.append((cur, tot)))
        self.assertEqual(rs[0]['status'], 'ok', rs[0].get('error'))
        self.assertTrue(pages_seen, "未收到页级进度回调")
        self.assertEqual(pages_seen[-1][0], pages_seen[-1][1])
        self.assertEqual(len(pages_seen), rs[0]['pages'])


class TestHybridMixedPdf(unittest.TestCase):
    """混合文档(文字页+图片页): 图片页必须OCR成可编辑文字, 文字页保留版式"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp(prefix="p2w_hybrid_")
        self.in_dir = os.path.join(self.tmp, "in")
        self.out_dir = os.path.join(self.tmp, "out")
        os.makedirs(self.in_dir)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_mixed_pdf_image_pages_become_editable_text(self):
        pdf = os.path.join(self.in_dir, "mixed.pdf")
        _make_mixed_pdf(pdf)
        c = PDFToWordConverter(self.in_dir, self.out_dir)
        rs = c.convert_files([pdf])
        self.assertEqual(len(rs), 1)
        r = rs[0]
        self.assertEqual(r['status'], 'ok', r.get('error'))
        self.assertEqual(r['method'], 'hybrid')
        self.assertEqual(r['pages'], 4)

        out_path = os.path.join(self.out_dir, r['output'])
        self.assertTrue(os.path.isfile(out_path))

        from docx import Document
        doc = Document(out_path)
        text = "\n".join(p.text for p in doc.paragraphs)

        # 文字页经 pdf2docx 保留
        self.assertIn("QUANTUM COVER PAGE", text)
        # 图片页 OCR 成可编辑文字
        self.assertIn("ALPHA", text)
        self.assertIn("BETA", text)
        # 仅第4页(纯色照片)触发"嵌入原图"兜底 => 恰好 1 张内嵌图片
        self.assertEqual(len(doc.inline_shapes), 1,
                         "图片页应被OCR成文字, 仅纯照片页保留1张原图")

        # 可编辑性: 修改段落后能保存
        non_empty = next(p for p in doc.paragraphs if p.text.strip())
        non_empty.text = "EDITED-" + non_empty.text
        doc.save(out_path)
        texts2 = [p.text for p in Document(out_path).paragraphs]
        self.assertTrue(any(t.startswith("EDITED-") for t in texts2))


class TestScannedPdfOcr(unittest.TestCase):
    """扫描件 OCR 路径: 输出必须是可编辑文字而非仅图片"""

    TRUMP_PDF = os.path.join(PROJECT_DIR, "特朗普.pdf")

    def setUp(self):
        if not os.path.isfile(self.TRUMP_PDF):
            self.skipTest("特朗普.pdf 不存在")
        self.out_dir = tempfile.mkdtemp(prefix="p2w_ocr_")

    def tearDown(self):
        shutil.rmtree(self.out_dir, ignore_errors=True)

    def test_scanned_pdf_produces_editable_text(self):
        c = PDFToWordConverter(PROJECT_DIR, self.out_dir)
        results = c.convert_files([self.TRUMP_PDF])
        self.assertEqual(len(results), 1)
        r = results[0]
        self.assertEqual(r['status'], 'ok', r.get('error'))
        self.assertEqual(r['method'], 'ocr')

        out_path = os.path.join(self.out_dir, r['output'])
        self.assertTrue(os.path.isfile(out_path))

        from docx import Document
        doc = Document(out_path)
        text = "\n".join(p.text for p in doc.paragraphs)
        # 必须有实质可编辑文字(非空段落足够多)
        non_empty = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertGreater(len(non_empty), 10, "OCR段落过少")
        self.assertGreater(len(text), 500, "可编辑文字过少")
        # 核心内容关键词(测试文件为CNN特朗普报道)
        self.assertIn("Trump", text)

        # 可编辑性: 修改段落后能保存
        non_empty_para = next(p for p in doc.paragraphs if p.text.strip())
        original = non_empty_para.text
        non_empty_para.text = "EDITED-" + original
        doc.save(out_path)
        doc2 = Document(out_path)
        texts2 = [p.text for p in doc2.paragraphs if p.text.strip()]
        self.assertTrue(any(t.startswith("EDITED-") for t in texts2),
                        "修改后保存失败, 文档不可编辑")


class TestIntegrationRealPdf(unittest.TestCase):
    """用 data/ 目录真实PDF做端到端验证(目录为空则跳过)"""

    def setUp(self):
        if not os.path.isdir(DATA_DIR):
            self.skipTest("data/ 目录不存在")
        self.pdfs = [f for f in os.listdir(DATA_DIR) if f.lower().endswith('.pdf')]
        if not self.pdfs:
            self.skipTest("data/ 目录中没有PDF")
        self.out_dir = tempfile.mkdtemp(prefix="p2w_e2e_")

    def tearDown(self):
        shutil.rmtree(self.out_dir, ignore_errors=True)

    def test_end_to_end_quality(self):
        c = PDFToWordConverter(DATA_DIR, self.out_dir)
        results = c.convert_all()

        self.assertEqual(len(results), len(self.pdfs))
        for r in results:
            self.assertIn(r['status'], ('ok',), f"{r['file']}: {r['error']}")
            self.assertGreater(r['pages'] or 0, 0)
            path = os.path.join(self.out_dir, r['output'])
            self.assertTrue(os.path.exists(path), f"未生成输出: {r['output']}")
            self.assertGreater(os.path.getsize(path), 100)


if __name__ == '__main__':
    unittest.main()
